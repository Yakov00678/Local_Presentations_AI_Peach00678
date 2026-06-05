import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def parse_markdown_to_slides(text):
    # Разделяем текст строго по тегу "# СЛАЙД"
    raw_slides = re.split(r'#\s*СЛАЙД\s*\d+', text)
    slides_data = []
    
    for raw in raw_slides:
        lines = [line.strip() for line in raw.split('\n') if line.strip()]
        if not lines:
            continue
            
        slide_info = {'title': '', 'content_lines': [], 'table': None}
        table_rows = []
        
        for line in lines:
            # Игнорируем разделители "---"
            if line.startswith('---'):
                continue
                
            if line.startswith('##'):
                slide_info['title'] = line.replace('##', '').strip()
            elif line.startswith('|'):
                # Пропускаем декоративную разделительную черту таблицы |---|---|
                if set(line.replace('|', '').replace('-', '').replace(' ', '')) <= {' '}:
                    continue
                cells = [c.strip() for c in line.split('|')][1:-1]
                if cells:
                    table_rows.append(cells)
            else:
                slide_info['content_lines'].append(line)
                
        if table_rows:
            slide_info['table'] = table_rows
            
        # Добавляем слайд, если в нем есть хоть какой-то полезный контент
        if slide_info['title'] or slide_info['content_lines'] or slide_info['table']:
            slides_data.append(slide_info)
        
    return slides_data

def create_presentation(slides_data):
    prs = Presentation()
    # Задаем современный широкоформатный размер 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    COLOR_TEXT = RGBColor(44, 62, 80)     
    COLOR_BLUE = RGBColor(41, 128, 185)   
    
    for data in slides_data:
        # Индекс 6 — это полностью пустой макет слайда в стандартном PowerPoint
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        
        # 1. Заголовок
        if data['title']:
            txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(1))
            tf = txBox.text_frame
            tf.word_wrap = True
            # ИСПРАВЛЕНО: берем конкретно первый абзац из коллекции
            p = tf.paragraphs[0]
            p.text = data['title']
            p.font.name = 'Trebuchet MS'
            p.font.size = Pt(38)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEXT
            
        # 2. Текст и списки
        if data['content_lines']:
            txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5))
            tf = txBox.text_frame
            tf.word_wrap = True
            
            for i, line in enumerate(data['content_lines']):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                
                if line.startswith('-'):
                    p.text = "  •  " + line.replace('-', '').strip()
                    p.font.size = Pt(20)
                    p.space_before = Pt(8)
                else:
                    p.text = line
                    p.font.size = Pt(22)
                    p.space_before = Pt(14)
                    
                p.font.name = 'Calibri'
                p.font.color.rgb = COLOR_TEXT
                    
        # 3. Монолитная таблица
        if data['table']:
            rows = len(data['table'])
            cols = max(len(row) for row in data['table']) if data['table'] else 0
            
            if rows > 0 and cols > 0:
                left = Inches(0.8)
                top = Inches(2.0)
                width = Inches(11.7)
                height = Inches(0.5 * rows)
                
                table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
                table = table_shape.table
                
                for r_idx, row in enumerate(data['table']):
                    for c_idx, val in enumerate(row):
                        if c_idx < len(table.columns):
                            cell = table.cell(r_idx, c_idx)
                            cell.text = val
                            
                            for p in cell.text_frame.paragraphs:
                                p.font.name = 'Calibri'
                                p.font.size = Pt(15)
                                p.font.color.rgb = COLOR_TEXT
                                
                                # Оформление шапки таблицы
                                if r_idx == 0:
                                    p.font.bold = True
                                    p.font.color.rgb = RGBColor(255, 255, 255)
                                    cell.fill.solid()
                                    cell.fill.fore_color.rgb = COLOR_BLUE

    # Создание папки и сохранение
    output_folder = "Presentations"
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)
        
    file_path = os.path.join(output_folder, "ready_presentation.pptx")
    prs.save(file_path)
    print(f"\n✅ Все ошибки устранены! Файл успешно сохранен: {file_path}")

# Исходный текст презентации
markdown_input = """
# СЛАЙД 2

## Интервью: анатомия диалога

Интервью — это не «просто поговорить». Это инструмент:
- Получения информации
- Формирования общественного мнения
- Фиксации исторического момента
- Раскрытия личности
- Анализа проблемы

Журналист = профессиональный посредник между героем и аудиторией

---

# СЛАЙД 3

## Типология интервью



| Тип | Цель | Характеристика |
|-----|------|----------------|
| Информационное | Получение фактов | Конкретные вопросы, фокус на событии |
| Портретное | Раскрытие личности | Открытые вопросы, акцент на биографии |
| Проблемное | Анализ явления | Вопросы на сравнение, гипотезы, экспертные оценки |
| Блиц | Быстрая фиксация позиции | Короткие вопросы, ограниченное время |

Выбор типа = задача редакции + аудитория + характер героя
"""

slides = parse_markdown_to_slides(markdown_input)
create_presentation(slides)
